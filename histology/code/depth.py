"""Coordinate extraction, depth projection, and rank normalization.

Parses stereology Excel files from the coordinates/ directory to extract
site XY positions, project onto a L2/3→L5/6 depth axis, and rank-normalize.
"""

import os
import re
import numpy as np
import pandas as pd
from scipy.stats import pearsonr, rankdata

from . import config


def extract_coordinates_from_excel(xlsx_path):
    """Parse a single stereology Excel file to extract site XY coordinates.

    Reads the `temp (2)` sheet:
    - Site mapping: rows 43-62, col H = Site, col J = X, col K = Y
    - Falls back to lookup table (cols A-C) + GridPos (col I) if X/Y are #N/A

    Parameters
    ----------
    xlsx_path : str or Path

    Returns
    -------
    dict : {site_number: (x, y)} for sites 1-20, or empty dict if unresolvable
    """
    import openpyxl
    wb = openpyxl.load_workbook(str(xlsx_path), data_only=True)
    ws = wb["temp (2)"]

    site_coords = {}
    for row in range(43, 63):
        site = ws.cell(row=row, column=8).value   # Column H = Site
        x = ws.cell(row=row, column=10).value      # Column J = X
        y = ws.cell(row=row, column=11).value      # Column K = Y

        if site is None:
            continue

        site_num = int(site)

        if x is not None and x != '#N/A' and y is not None and y != '#N/A':
            site_coords[site_num] = (float(x), float(y))
        else:
            # Try lookup fallback: GridPos in col I → lookup in cols A-C
            grid_pos = ws.cell(row=row, column=9).value
            if grid_pos is not None:
                gp = int(grid_pos)
                # Search lookup table
                for lrow in range(3, 200):
                    num = ws.cell(row=lrow, column=1).value
                    if num is None:
                        break
                    if int(num) == gp:
                        lx = ws.cell(row=lrow, column=2).value
                        ly = ws.cell(row=lrow, column=3).value
                        if lx is not None and ly is not None:
                            site_coords[site_num] = (float(lx), float(ly))
                        break

    wb.close()
    return site_coords


def batch_extract_coordinates(coord_dir=None):
    """Extract site coordinates from all Excel files in a directory.

    Parameters
    ----------
    coord_dir : str or Path, optional
        Directory containing {subject} - {L/R}.xlsx files.
        Defaults to config.PROJECT_DIR / "coordinates".

    Returns
    -------
    DataFrame with columns: subject, section, site, x, y
    """
    coord_dir = coord_dir or config.PROJECT_DIR / "coordinates"

    files = sorted([
        f for f in os.listdir(coord_dir)
        if f.endswith('.xlsx') and not f.startswith('~')
    ])

    all_rows = []
    failed = []

    for f in files:
        m = re.match(r'(\d+)\s*-\s*([LRlr])\.xlsx', f)
        if not m:
            failed.append(f"Can't parse: {f}")
            continue

        subj_id = m.group(1)
        section = m.group(2).upper()

        try:
            site_coords = extract_coordinates_from_excel(os.path.join(coord_dir, f))
            for site_num, (x, y) in site_coords.items():
                all_rows.append({
                    'subject': subj_id,
                    'section': section,
                    'site': site_num,
                    'x': x,
                    'y': y,
                })
        except Exception as e:
            failed.append(f"{f}: {e}")

    if failed:
        print(f"Coordinate extraction warnings: {failed}")

    df = pd.DataFrame(all_rows)

    # Apply subject ID fixes
    df['subject'] = df['subject'].map(lambda x: config.ID_FIXES.get(x, x))

    print(f"Extracted coordinates: {len(df)} sites from "
          f"{df.groupby(['subject', 'section']).ngroups} subject-sections "
          f"({df['subject'].nunique()} subjects)")
    return df


def compute_depth_for_section(section_df):
    """Compute depth projection and rank normalization for one subject-section.

    Projects onto L2/3-centroid → L5/6-centroid axis, then rank-normalizes.

    Parameters
    ----------
    section_df : DataFrame
        Must have columns: site, x, y. Sites 1-10 = L2/3, 11-20 = L5/6.

    Returns
    -------
    DataFrame with added columns: depth_raw, depth
    """
    l23 = section_df[section_df['site'] <= 10][['x', 'y']].values
    l56 = section_df[section_df['site'] > 10][['x', 'y']].values

    if len(l23) == 0 or len(l56) == 0:
        section_df = section_df.copy()
        section_df['depth_raw'] = np.nan
        section_df['depth'] = np.nan
        return section_df

    l23_centroid = l23.mean(axis=0)
    l56_centroid = l56.mean(axis=0)
    axis = l56_centroid - l23_centroid
    axis_len = np.linalg.norm(axis)

    if axis_len < 1e-6:
        section_df = section_df.copy()
        section_df['depth_raw'] = np.nan
        section_df['depth'] = np.nan
        return section_df

    axis_unit = axis / axis_len

    projections = []
    for _, row in section_df.iterrows():
        vec = np.array([row['x'], row['y']]) - l23_centroid
        projections.append(np.dot(vec, axis_unit))

    section_df = section_df.copy()
    section_df['depth_raw'] = projections

    # Rank normalize to [0, 1]
    ranks = rankdata(projections)
    n = len(projections)
    section_df['depth'] = (ranks - 1) / (n - 1)

    return section_df


def compute_all_depths(coords_df):
    """Compute depth for all subject-sections.

    Parameters
    ----------
    coords_df : DataFrame
        Output of batch_extract_coordinates().

    Returns
    -------
    DataFrame with added columns: depth_raw, depth
    """
    results = []
    for (subj, section), grp in coords_df.groupby(['subject', 'section']):
        result = compute_depth_for_section(grp)
        results.append(result)

    depth_df = pd.concat(results, ignore_index=True)

    # Validate: depth should correlate with binary layer
    layer_binary = (depth_df['site'] > 10).astype(int)
    r, p = pearsonr(depth_df['depth'].dropna(), layer_binary[depth_df['depth'].notna()])
    print(f"Depth validation: r={r:.3f} vs binary layer (p={p:.2e})")

    return depth_df


def validate_depth_with_vip(df, depth_col="depth"):
    """Compute per-subject VIP-depth correlation to validate depth assignment.

    VIP is concentrated in superficial layers, so a correct depth axis
    should produce a negative VIP-depth correlation.

    Parameters
    ----------
    df : DataFrame
        Must have columns: subject, VIP, and depth_col.

    Returns
    -------
    DataFrame with columns: subject, vip_depth_r, vip_depth_p, quality
        quality: 'strong' (r < -0.3), 'acceptable' (-0.3 ≤ r < 0), 'failed' (r ≥ 0)
    """
    results = []
    for subj, grp in df.groupby("subject"):
        vip_data = grp.dropna(subset=["VIP", depth_col])
        if len(vip_data) < 5 or vip_data["VIP"].std() < 1e-6:
            results.append({"subject": subj, "vip_depth_r": np.nan,
                            "vip_depth_p": np.nan, "quality": "failed"})
            continue

        r, p = pearsonr(vip_data[depth_col], vip_data["VIP"])

        if r < -0.3:
            quality = "strong"
        elif r < 0:
            quality = "acceptable"
        else:
            quality = "failed"

        results.append({"subject": subj, "vip_depth_r": r,
                        "vip_depth_p": p, "quality": quality})

    return pd.DataFrame(results)
